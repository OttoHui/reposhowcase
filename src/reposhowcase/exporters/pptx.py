from __future__ import annotations

from pathlib import Path

from ..models import Storyboard


def export_pptx(storyboard: Storyboard, output: str | Path) -> Path:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX export requires `pip install reposhowcase[pptx]`.") from exc
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    for slide_data in storyboard.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_data.title
        frame = slide.placeholders[1].text_frame
        frame.text = slide_data.body
        for bullet in slide_data.bullets:
            paragraph = frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
        notes = slide.notes_slide.notes_text_frame
        notes.text = "Evidence: " + ", ".join(item.path for item in slide_data.evidence)
    presentation.save(destination)
    return destination
